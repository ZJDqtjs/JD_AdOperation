# -*- coding: utf-8 -*-
"""导出「数据看板 → 商品数据」每个 SKU 的付款佣金金额到 Excel。

数据来源（复用已登录 Profile，页面内 fetch）：
  POST https://atoms-api.jd.com/union/dataBoard/chooseGoods  商品列表（含 SKU/名称/下单金额）
  POST https://atoms-api.jd.com/union/dataBoard/orderData    单个 SKU 的 CPS 收支（含 payCommAmount）
  POST https://atoms-api.jd.com/union/dataBoard/targetCompare 分日趋势（账户级合计）

用法（容器内）：
    docker cp scripts/export_paycomm_by_sku.py jd-roi-dev:/tmp/ && \
    docker exec -w /app -e PYTHONPATH=/app jd-roi-dev \
      python /tmp/export_paycomm_by_sku.py 2026-08-01 2026-08-31 main

产物：/data/reports/paycomm_<start>_<end>_<account>.xlsx
      （宿主机 docker-dev/data/reports/ 下）
"""
from __future__ import annotations

import json
import random
import sys
import time
from pathlib import Path

from jd_roi import settings
from jd_roi.browser import launch_profile

BASE = "https://atoms-api.jd.com/union/dataBoard"
PAGE = "https://jzt.jd.com/jtk/#/data-board"

_JS = """
async ({url, payload}) => {
  const r = await fetch(url, {method:'POST', credentials:'include',
    headers:{'Content-Type':'application/json',
             'accept':'application/json, text/plain, */*',
             'referer':'https://jzt.jd.com/jtk/',
             'language':'zh_CN','siteid':'0','loginmode':'0'},
    body: JSON.stringify(payload)});
  return {status: r.status, text: await r.text()};
}
"""

# 限速：单次请求之间留随机间隔，避免请求过快触发风控（掉登录态要重扫码）
DELAY_SKU = (0.8, 1.6)
DELAY_PAGE = (0.6, 1.2)


def _sleep(rng: tuple[float, float]) -> None:
    time.sleep(random.uniform(*rng))


# 页面字段 -> Excel 列名（顺序即列顺序）
FIELDS = [
    ("payCommAmount", "付款佣金金额(元)"),
    ("payCommRate", "付款佣金比例(%)"),
    ("preCommAmount", "预估佣金金额(元)"),
    ("preCommRate", "预估佣金比例(%)"),
    ("clicks", "点击数"),
    ("uv", "引入UV"),
    ("orderCnt", "下单订单量"),
    ("orderCvr", "下单流量转化率(%)"),
    ("orderGmv", "下单订单金额(元)"),
    ("payOrderCnt", "付款订单量"),
    ("fiOrderCnt", "完成订单量"),
    ("fiOrderGmv", "完成订单金额(元)"),
    ("orderUserCnt", "下单用户数"),
    ("roi", "ROI"),
]


def post(page, path: str, payload: dict, retries: int = 2) -> dict:
    url = f"{BASE}/{path}?requestFrom=0"
    last = ""
    for _ in range(retries + 1):
        try:
            res = page.evaluate(_JS, {"url": url, "payload": payload})
            if res["status"] != 200:
                last = f"HTTP {res['status']}"
            else:
                data = json.loads(res["text"])
                if str(data.get("code")) == "1":
                    return data
                last = json.dumps(data, ensure_ascii=False)[:200]
        except Exception as exc:  # noqa: BLE001
            last = f"{type(exc).__name__}: {exc}"
        time.sleep(1.5)
    raise RuntimeError(f"{path} 失败: {last}")


def list_goods(page, start: str, end: str, page_size: int = 100) -> list[dict]:
    rows: list[dict] = []
    idx = 1
    while True:
        d = post(page, "chooseGoods", {"displayType": 1, "itemId": "", "pageIndex": idx,
                                      "pageSize": page_size, "startDate": start, "endDate": end})
        data = d.get("data") or {}
        batch = data.get("data") or []
        rows.extend(batch)
        total = int(data.get("total") or 0)
        if not batch or len(rows) >= total or idx > 50:
            break
        idx += 1
        _sleep(DELAY_PAGE)
    return rows


def sku_order_data(page, sku: int, start: str, end: str) -> dict:
    d = post(page, "orderData", {"startDate": start, "endDate": end, "realTime": False,
                                 "skuId": sku, "spuId": None})
    return d.get("data") or {}


def daily_total(page, start: str, end: str) -> list[dict]:
    d = post(page, "targetCompare", {"startDate": start, "endDate": end,
                                     "realTime": False, "skuId": None, "spuId": None})
    return d.get("data") or []


def write_xlsx(path: Path, start: str, end: str, rows: list[dict], daily: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "SKU付款佣金"
    head = ["SKU", "商品名称"] + [name for _, name in FIELDS]
    ws.append(head)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([r["skuId"], r["itemName"]] + [r.get(k) for k, _ in FIELDS])
    total_pay = round(sum((r.get("payCommAmount") or 0) for r in rows), 2)
    ws.append([])
    ws.append(["合计", f"{len(rows)} 个 SKU", total_pay])
    ws.freeze_panes = "C2"
    widths = [15, 46] + [17] * len(FIELDS)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[2:]:
            if isinstance(c.value, (int, float)):
                c.number_format = "#,##0.00" if isinstance(c.value, float) else "#,##0"

    ws2 = wb.create_sheet("分日合计")
    ws2.append(["日期", "付款佣金金额(元)", "预估佣金金额(元)", "点击数", "下单订单量",
                "下单订单金额(元)", "完成订单量", "完成订单金额(元)", "ROI"])
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = fill
    for d in daily:
        ws2.append([d.get("date"), d.get("payCommAmount"), d.get("preCommAmount"), d.get("clicks"),
                    d.get("orderCnt"), d.get("orderGmv"), d.get("fiOrderCnt"),
                    d.get("fiOrderGmv"), d.get("roi")])
    ws2.append([])
    ws2.append(["合计", round(sum((d.get("payCommAmount") or 0) for d in daily), 2),
                round(sum((d.get("preCommAmount") or 0) for d in daily), 2),
                sum((d.get("clicks") or 0) for d in daily),
                sum((d.get("orderCnt") or 0) for d in daily),
                round(sum((d.get("orderGmv") or 0) for d in daily), 2),
                sum((d.get("fiOrderCnt") or 0) for d in daily),
                round(sum((d.get("fiOrderGmv") or 0) for d in daily), 2), None])
    for i, w in enumerate([13, 18, 18, 12, 13, 18, 13, 18, 10], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "B2"

    ws3 = wb.create_sheet("说明")
    ws3.append(["统计区间", f"{start} ~ {end}"])
    ws3.append(["数据来源", "京准通 数据看板 → 商品数据（atoms-api /union/dataBoard）"])
    ws3.append(["导出时间", time.strftime("%Y-%m-%d %H:%M:%S")])
    ws3.append(["口径", "每个 SKU 单独请求 orderData；付款佣金金额=payCommAmount（结算/付款口径，与页面红框一致）"])
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 90

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def main() -> int:
    args = [a for a in sys.argv[1:] if a]
    start = args[0] if len(args) > 0 else "2026-08-01"
    end = args[1] if len(args) > 1 else "2026-08-31"
    account = args[2] if len(args) > 2 else "main"

    out = Path(settings.DATA_DIR) / "reports" / f"paycomm_{start}_{end}_{account}.xlsx"
    raw = Path(settings.DATA_DIR) / "reports" / f"paycomm_{start}_{end}_{account}.json"

    pw, context = launch_profile(headless=True, account=account)
    page = context.new_page()
    try:
        page.goto(PAGE, wait_until="domcontentloaded", timeout=90000)
        page.wait_for_timeout(12000)

        goods = list_goods(page, start, end)
        print(f"[goods] {len(goods)} 个 SKU（{start}~{end}）")

        rows: list[dict] = []
        for i, g in enumerate(goods, 1):
            sku = int(g["itemId"])
            try:
                od = sku_order_data(page, sku, start, end)
            except Exception as exc:  # noqa: BLE001
                print(f"  [fail] {sku}: {exc}")
                od = {}
            rec = {"skuId": sku, "itemName": g.get("itemName")}
            for k, _ in FIELDS:
                rec[k] = od.get(k)
            rows.append(rec)
            print(f"  [{i}/{len(goods)}] {sku} payCommAmount={rec.get('payCommAmount')}")
            _sleep(DELAY_SKU)

        try:
            daily = daily_total(page, start, end)
        except Exception:  # noqa: BLE001
            daily = []

        rows.sort(key=lambda r: (r.get("payCommAmount") or -1), reverse=True)
        write_xlsx(out, start, end, rows, daily)
        raw.write_text(json.dumps({"start": start, "end": end, "rows": rows, "daily": daily},
                                 ensure_ascii=False, indent=1), encoding="utf-8")
        total = sum((r.get("payCommAmount") or 0) for r in rows)
        print(f"\n[OK] {len(rows)} 个 SKU，付款佣金合计 {total:,.2f} -> {out}")
        return 0
    finally:
        context.close()
        pw.stop()


if __name__ == "__main__":
    sys.exit(main())
